# Nombre: Valentina Porras Alfaro
#Nombre: Jose Andre Mora Paredes
import os
import csv
from colorama import Fore, Style, init
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from procesamiento import limpiar_y_filtrar_lecturas, calcular_metricas_produccion

#Cap 9: Inicializacion de colorama
init(autoreset=True)

#Cap 11: Creacion de rytas y directorios
ruta_base = os.path.dirname(os.path.abspath(__file__))
ruta_datos = os.path.join(ruta_base, "Datos_Planta")
os.makedirs(ruta_datos, exist_ok=True)

#FASE 2: Procesamiento de datos

#Datos crudos de produccion
datos_crudos = [
    "23.5 \n", "ERROR_ZONA_3", "24.8\r\n", "nan ", "22.1"
]

#Limpieza y filtrado de lecturas
temperaturas, incidencias = limpiar_y_filtrar_lecturas(datos_crudos)
 
#Calculo de metricas 
minimo, maximo, promedio = calcular_metricas_produccion(temperaturas)

# Mostrar resultados 
print(Fore.CYAN + Style.BRIGHT + "REPORTE DE PRODUCCION")
print(Fore.GREEN + f"Temperaturas validas: {temperaturas}")
print(Fore.YELLOW + f"Temperatura minima: {minimo}")
print(Fore.YELLOW + f"Temperatura maxima: {maximo}")
print(Fore.YELLOW + f"Temperatura promedio: {promedio:.2f}")

#Guardar incidencias en archivo de texto
ruta_bitacora = os.path.join(ruta_datos, "bitacora_alertas.txt")

with open(ruta_bitacora, "a", encoding= "utf-8") as archivo:
    for inicidencia in incidencias: 
        archivo.write(inicidencia + "\n")

print(Fore.RED + f"Incidencias registradas: {len(incidencias)}")

# Exportar metricas a CSV
ruta_csv = os.path.join(ruta_datos, "reporte_metricas.csv")
with open(ruta_csv, "w", newline="", encoding="utf-8") as archivo:
    escritor = csv.writer(archivo, delimiter=";")
    escritor.writerow(["Métrica", "Valor"])
    escritor.writerow(["Mínimo", minimo])
    escritor.writerow(["Máximo", maximo])
    escritor.writerow(["Promedio", promedio])

print(Fore.GREEN + "Reporte CSV generado correctamente.")

# FASE 3: Automatización de Excel

# Crear libro de Excel
libro = Workbook()
hoja = libro.active
hoja.title = "Producción_Mensual"

# Encabezados
hoja["A1"] = "Código Componente"
hoja["B1"] = "Cantidad Producida"

# Formato de encabezados
for celda in hoja[1]:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill(
        fill_type="solid",
        fgColor="000080"
    )

# Datos simulados de producción
datos_produccion = [
    ["COMP-001", 150],
    ["COMP-002", 200],
    ["COMP-003", 175],
    ["COMP-004", 220],
    ["COMP-005", 190]
]

# Agregar registros
for registro in datos_produccion:
    hoja.append(registro)

# Guardar archivo Excel
ruta_excel = os.path.join(ruta_datos, "Reporte_Calidad.xlsx")
libro.save(ruta_excel)

print(Fore.GREEN + "Reporte Excel generado correctamente.")

# Bloque de auditoría
libro = load_workbook(ruta_excel)
hoja = libro["Producción_Mensual"]

valor_anterior = hoja["B2"].value

# Corrección de cantidad producida
hoja["B2"] = 160

# Guardar corrección
libro.save(ruta_excel)

print(Fore.CYAN + "Auditoría de calidad realizada.")
print(Fore.YELLOW + f"Valor anterior: {valor_anterior}")
print(Fore.GREEN + f"Valor corregido: {hoja['B2'].value}")

print(Fore.GREEN + Style.BRIGHT + "Proceso finalizado correctamente.")










